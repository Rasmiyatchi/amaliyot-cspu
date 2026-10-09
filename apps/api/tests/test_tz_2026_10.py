"""TZ 08.10.2026: audit diff/eksport, qayta biriktirish va ommaviy sxemalar, qurilma uzish."""

from datetime import UTC, date, datetime
from uuid import uuid4

import pytest
from app.models.audit_log import AuditLog
from app.models.enums import Semester
from app.schemas.assignment_bulk import (
    AssignmentBulkChanges,
    AssignmentBulkUpdateRequest,
    ReassignRequest,
    ReassignScope,
    ReassignTarget,
)
from app.schemas.broadcast import BroadcastCreate
from app.schemas.student import DeviceResetRequest
from app.services.audit_log import (
    EXPORT_MAX_ROWS,
    _affected,
    _format_changes,
    _jsonable,
    build_audit_xlsx,
    diff,
)
from pydantic import ValidationError


class TestAuditDiff:
    def test_faqat_ozgargan_maydonlar(self) -> None:
        sid = uuid4()
        before = {"supervisor_id": sid, "start_date": date(2026, 2, 1), "notes": None}
        after = {"supervisor_id": sid, "start_date": date(2026, 3, 1), "notes": "x"}
        out = diff(before, after)
        assert [d["field"] for d in out] == ["start_date", "notes"]
        assert out[0]["before"] == "2026-02-01" and out[0]["after"] == "2026-03-01"

    def test_jsonable_uuid_enum_sana(self) -> None:
        uid = uuid4()
        out = _jsonable({"a": uid, "b": Semester.FALL, "c": [date(2026, 1, 2)], "d": {1: None}})
        assert out == {"a": str(uid), "b": "fall", "c": ["2026-01-02"], "d": {"1": None}}

    def test_fields_bilan_cheklash(self) -> None:
        out = diff({"a": 1, "b": 2}, {"a": 9, "b": 9}, fields=["a"])
        assert out == [{"field": "a", "before": 1, "after": 9}]


class TestAuditExport:
    def _row(self, **kw: object) -> AuditLog:
        base: dict[str, object] = {
            "actor_user_id": None,
            "actor_role": "super_admin",
            "actor_name": "Admin A.",
            "action": "update",
            "entity_type": "practice_assignment",
            "entity_id": uuid4(),
            "summary": "Biriktirish tahrirlandi",
            "metadata_json": None,
            "ip": "10.0.0.1",
            "user_agent": "ua",
        }
        base.update(kw)
        row = AuditLog(**base)
        row.created_at = datetime(2026, 10, 9, 5, 0, tzinfo=UTC)  # 10:00 Toshkent
        return row

    def test_xlsx_sarlavha_va_qatorlar(self) -> None:
        from io import BytesIO

        from openpyxl import load_workbook

        rows = [
            self._row(
                metadata_json={
                    "changes": [
                        {"field": "end_date", "before": "2026-05-01", "after": "2026-06-01"}
                    ]
                }
            ),
            self._row(
                action="reassign",
                summary="Qayta biriktirish",
                metadata_json={"affected_count": 25, "result": "partial"},
            ),
        ]
        wb = load_workbook(BytesIO(build_audit_xlsx(rows)))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "№"
        assert ws.cell(row=2, column=2).value == "09.10.2026 10:00:00"  # Toshkent vaqti
        assert ws.cell(row=2, column=5).value == "update"
        assert "end_date: 2026-05-01 → 2026-06-01" in ws.cell(row=2, column=9).value
        assert ws.cell(row=3, column=10).value == "25 · partial"
        assert ws.max_row == 3

    def test_format_helpers(self) -> None:
        assert _format_changes(None) == ""
        assert _affected({"affected_count": 3}) == "3"
        assert _affected({"result": "ok"}) == "ok"
        txt = _format_changes({"before": {"x": 1}, "after": {"x": 2}})
        assert txt == "x: 1 → 2"
        assert EXPORT_MAX_ROWS == 10_000


class TestReassignSchemas:
    def test_qamrov_shart(self) -> None:
        with pytest.raises(ValidationError, match="Qamrov tanlanmagan"):
            ReassignScope(academic_year_id=uuid4())

    def test_sana_tartibi(self) -> None:
        with pytest.raises(ValidationError, match="start_date"):
            ReassignTarget(start_date=date(2026, 3, 1), end_date=date(2026, 2, 1))

    def test_default_dry_run_va_hafta_kunlari(self) -> None:
        req = ReassignRequest(
            source=ReassignScope(academic_year_id=uuid4(), group_id=uuid4(), semester="fall"),
            target=ReassignTarget(
                semester="spring",
                start_date=date(2026, 2, 2),
                end_date=date(2026, 5, 30),
                required_weekdays=[3, 1, 1],
            ),
        )
        assert req.dry_run is True
        assert req.target.required_weekdays == [1, 3]
        assert req.target.keep_weekdays is True and req.target.activate is False

    def test_hafta_kuni_oraliq(self) -> None:
        with pytest.raises(ValidationError):
            ReassignTarget(
                start_date=date(2026, 2, 2), end_date=date(2026, 5, 30), required_weekdays=[0]
            )


class TestBulkUpdateSchemas:
    def test_ozgarish_shart(self) -> None:
        with pytest.raises(ValidationError, match="maydon tanlanmagan"):
            AssignmentBulkUpdateRequest(assignment_ids=[uuid4()], changes=AssignmentBulkChanges())

    def test_qamrov_shart(self) -> None:
        with pytest.raises(ValidationError, match="Qamrov tanlanmagan"):
            AssignmentBulkUpdateRequest(changes=AssignmentBulkChanges(required_weekdays=[1, 2]))

    def test_supervizorni_olib_tashlash_aniq_null(self) -> None:
        req = AssignmentBulkUpdateRequest(
            group_id=uuid4(), changes=AssignmentBulkChanges(supervisor_id=None)
        )
        # exclude_unset: null yuborilgan → "olib tashlash", yuborilmagan → tegilmaydi
        assert req.changes.model_dump(exclude_unset=True) == {"supervisor_id": None}
        assert "required_weekdays" not in req.changes.model_dump(exclude_unset=True)


class TestBroadcastSchema:
    def test_auditoriya_nishoni(self) -> None:
        with pytest.raises(ValidationError, match="Guruh tanlanmagan"):
            BroadcastCreate(audience="group", subject="Mavzu", body="Matn matn")
        with pytest.raises(ValidationError, match="Fakultet tanlanmagan"):
            BroadcastCreate(audience="faculty", subject="Mavzu", body="Matn matn")
        with pytest.raises(ValidationError, match="Talabalar tanlanmagan"):
            BroadcastCreate(audience="students", subject="Mavzu", body="Matn matn")
        ok = BroadcastCreate(audience="all_students", subject="Mavzu", body="Matn matn")
        assert ok.dry_run is False

    def test_mavzu_uzunligi(self) -> None:
        with pytest.raises(ValidationError):
            BroadcastCreate(audience="everyone", subject="ab", body="Matn matn")


class TestDeviceResetSchema:
    def test_tasdiqsiz_qollash_mumkin_emas(self) -> None:
        with pytest.raises(ValidationError, match="tasdiqlang"):
            DeviceResetRequest(scope="all", dry_run=False)
        ok = DeviceResetRequest(scope="all", dry_run=False, confirm=True)
        assert ok.scope == "all"

    def test_nishon(self) -> None:
        with pytest.raises(ValidationError, match="Guruh tanlanmagan"):
            DeviceResetRequest(scope="group")
        assert DeviceResetRequest(scope="students", student_ids=[uuid4()]).dry_run is True
