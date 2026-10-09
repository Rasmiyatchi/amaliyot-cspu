"""Audit log service — yozish + ro'yxat olish + Excel eksport.

Jurnal append-only: DB trigger (`trg_audit_logs_append_only`) UPDATE/DELETE ni rad etadi,
API'da ham tahrirlash/o'chirish endpointlari yo'q. Ommaviy amallar bitta yozuv bilan
(`affected_count`, `result`) yoziladi; tahrirlashlar `changes: [{field, before, after}]`
ko'rinishida eski va yangi qiymatlarni saqlaydi.
"""

from collections.abc import Iterable
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any
from uuid import UUID

from fastapi import Request
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.clock import UZB_TZ
from app.core.request_meta import client_ip
from app.models.academic import Direction, Group
from app.models.audit_log import AuditLog
from app.models.student import Student
from app.models.supervisor import Supervisor
from app.models.user import User
from app.services.search_utils import like_pattern, normalized_col

#: Excel eksportidagi maksimal qatorlar — undan ko'pi filtr bilan toraytiriladi
EXPORT_MAX_ROWS = 10_000
#: Ommaviy amal metadata'sida saqlanadigan maksimal element (diff) soni
BULK_DETAIL_LIMIT = 300


def _full_name(user: User) -> str:
    parts = [user.last_name or "", user.first_name or ""]
    return " ".join(p for p in parts if p).strip() or user.username


def _jsonable(value: Any) -> Any:
    """UUID/sana/enum qiymatlarini JSON'ga yoziladigan ko'rinishga keltiradi."""
    if value is None or isinstance(value, bool | int | float | str):
        return value
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    if hasattr(value, "value"):  # StrEnum
        return value.value
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, list | tuple | set):
        return [_jsonable(v) for v in value]
    return str(value)


def diff(
    before: dict[str, Any], after: dict[str, Any], fields: Iterable[str] | None = None
) -> list[dict[str, Any]]:
    """Eski va yangi qiymatlar farqi — faqat o'zgargan maydonlar.

    `fields` berilsa shu maydonlar tekshiriladi, aks holda `after` kalitlari.
    """
    keys = list(fields) if fields is not None else list(after.keys())
    out: list[dict[str, Any]] = []
    for key in keys:
        old = _jsonable(before.get(key))
        new = _jsonable(after.get(key))
        if old != new:
            out.append({"field": key, "before": old, "after": new})
    return out


async def log(
    db: AsyncSession,
    *,
    actor: User | None,
    action: str,
    entity_type: str,
    entity_id: UUID | None,
    summary: str,
    metadata: dict[str, Any] | None = None,
    request: Request | None = None,
) -> None:
    """Audit yozuvi qo'shadi. Hech qachon transaction'ni rollback qilmaydi.

    Asosiy commit'dan ALOHIDA flush qilinadi — caller commit qilishi kerak.
    Agar audit yozish o'zi xato bersa — silently log qilinadi (asosiy oqim
    yiqilmasligi uchun).
    """
    try:
        ip = None
        ua = None
        if request:
            ip = client_ip(request)
            ua = request.headers.get("user-agent")

        entry = AuditLog(
            actor_user_id=actor.id if actor else None,
            actor_role=actor.role.value if actor else None,
            actor_name=_full_name(actor) if actor else None,
            action=action,
            entity_type=entity_type,
            entity_id=entity_id,
            summary=summary[:500],
            metadata_json=_jsonable(metadata) if metadata is not None else None,
            ip=ip,
            user_agent=ua,
        )
        db.add(entry)
        await db.flush()
    except Exception:
        # Audit log xatosi asosiy oqimni to'xtatmasin
        from loguru import logger

        logger.exception("Audit log yozishda xato")


async def log_committed(
    db: AsyncSession,
    *,
    actor: User | None,
    action: str,
    entity_type: str,
    entity_id: UUID | None,
    summary: str,
    metadata: dict[str, Any] | None = None,
    request: Request | None = None,
) -> None:
    """Asosiy oqim xato bilan tugaganda ham (masalan, muvaffaqiyatsiz login) yozuv
    saqlanishi uchun — o'zi commit qiladi; xato bo'lsa jim rollback."""
    try:
        await db.rollback()
        await log(
            db,
            actor=actor,
            action=action,
            entity_type=entity_type,
            entity_id=entity_id,
            summary=summary,
            metadata=metadata,
            request=request,
        )
        await db.commit()
    except Exception:
        from loguru import logger

        logger.exception("Audit log (committed) yozishda xato")
        try:
            await db.rollback()
        except Exception:  # noqa: BLE001
            logger.warning("Audit log rollback ham xato berdi")


def _actor_faculty_subquery(faculty_id: UUID) -> Any:
    """Fakultet bo'yicha filtr: admin (users.faculty_id), talaba (guruh → yo'nalish),
    supervizor (supervisors.faculty_id) — shu fakultetga tegishli foydalanuvchilar."""
    admins = select(User.id).where(User.faculty_id == faculty_id)
    students = (
        select(Student.user_id)
        .join(Group, Group.id == Student.group_id)
        .join(Direction, Direction.id == Group.direction_id)
        .where(Direction.faculty_id == faculty_id)
    )
    supervisors = select(Supervisor.user_id).where(Supervisor.faculty_id == faculty_id)
    return admins.union(students, supervisors)


def _filters(
    *,
    actor_user_id: UUID | None,
    action: str | None,
    entity_type: str | None,
    date_from: date | None,
    date_to: date | None,
    search: str | None,
    faculty_id: UUID | None,
    exclude_actions: list[str] | None = None,
) -> list[Any]:
    conds: list[Any] = []
    if exclude_actions:
        conds.append(AuditLog.action.not_in(exclude_actions))
    if actor_user_id:
        conds.append(AuditLog.actor_user_id == actor_user_id)
    if action:
        conds.append(AuditLog.action == action)
    if entity_type:
        conds.append(AuditLog.entity_type == entity_type)
    # Sanalar Toshkent kuni bo'yicha: [date_from 00:00, date_to 24:00)
    if date_from:
        conds.append(
            AuditLog.created_at >= datetime.combine(date_from, datetime.min.time(), UZB_TZ)
        )
    if date_to:
        conds.append(
            AuditLog.created_at
            < datetime.combine(date_to + timedelta(days=1), datetime.min.time(), UZB_TZ)
        )
    if search and search.strip():
        pattern = like_pattern(search)
        conds.append(
            or_(
                normalized_col(AuditLog.summary).like(pattern, escape="\\"),
                normalized_col(func.coalesce(AuditLog.actor_name, "")).like(pattern, escape="\\"),
            )
        )
    if faculty_id:
        conds.append(AuditLog.actor_user_id.in_(_actor_faculty_subquery(faculty_id)))
    return conds


async def list_logs(
    db: AsyncSession,
    *,
    offset: int = 0,
    limit: int = 50,
    actor_user_id: UUID | None = None,
    action: str | None = None,
    entity_type: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    search: str | None = None,
    faculty_id: UUID | None = None,
    exclude_actions: list[str] | None = None,
) -> tuple[list[AuditLog], int]:
    conds = _filters(
        actor_user_id=actor_user_id,
        action=action,
        entity_type=entity_type,
        date_from=date_from,
        date_to=date_to,
        search=search,
        faculty_id=faculty_id,
        exclude_actions=exclude_actions,
    )
    total = (await db.execute(select(func.count(AuditLog.id)).where(*conds))).scalar_one()
    rows = (
        (
            await db.execute(
                select(AuditLog)
                .where(*conds)
                .order_by(AuditLog.created_at.desc())
                .offset(offset)
                .limit(limit)
            )
        )
        .scalars()
        .all()
    )
    return list(rows), total


async def list_actions(db: AsyncSession) -> list[str]:
    """Filtr ro'yxati uchun jurnalda uchragan amal turlari."""
    rows = (await db.execute(select(AuditLog.action).distinct().order_by(AuditLog.action))).all()
    return [r[0] for r in rows]


# ─── Excel eksport ────────────────────────────────────────

_EXPORT_HEADERS = [
    "№",
    "Vaqt (Toshkent)",
    "Kim",
    "Rol",
    "Amal",
    "Obyekt turi",
    "Obyekt ID",
    "Tafsilot",
    "O'zgarishlar (eski → yangi)",
    "Ta'sir (soni)",
    "IP",
]


def _format_changes(meta: dict[str, Any] | None) -> str:
    if not meta:
        return ""
    parts: list[str] = []
    changes = meta.get("changes")
    if isinstance(changes, list):
        for ch in changes[:20]:
            if isinstance(ch, dict):
                parts.append(f"{ch.get('field')}: {ch.get('before')} → {ch.get('after')}")
    before, after = meta.get("before"), meta.get("after")
    if isinstance(before, dict) and isinstance(after, dict):
        for d in diff(before, after)[:20]:
            parts.append(f"{d['field']}: {d['before']} → {d['after']}")
    if not parts:
        # Boshqa metadata — qisqa JSON ko'rinishda
        import json

        small = {k: v for k, v in meta.items() if k not in ("items", "changes")}
        text = json.dumps(small, ensure_ascii=False)
        return text[:500]
    return "\n".join(parts)


def _affected(meta: dict[str, Any] | None) -> str:
    if not meta:
        return ""
    n = meta.get("affected_count")
    result = meta.get("result")
    if n is None and result is None:
        return ""
    return f"{n if n is not None else ''}{' · ' + str(result) if result else ''}".strip(" ·")


def build_audit_xlsx(rows: list[AuditLog]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit jurnali"
    fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    for i, h in enumerate(_EXPORT_HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fill
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, r in enumerate(rows, start=1):
        row = idx + 1
        created = r.created_at.astimezone(UZB_TZ).strftime("%d.%m.%Y %H:%M:%S")
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=created)
        ws.cell(row=row, column=3, value=r.actor_name or "")
        ws.cell(row=row, column=4, value=r.actor_role or "")
        ws.cell(row=row, column=5, value=r.action)
        ws.cell(row=row, column=6, value=r.entity_type)
        ws.cell(row=row, column=7, value=str(r.entity_id) if r.entity_id else "")
        ws.cell(row=row, column=8, value=r.summary)
        ws.cell(row=row, column=9, value=_format_changes(r.metadata_json))
        ws.cell(row=row, column=10, value=_affected(r.metadata_json))
        ws.cell(row=row, column=11, value=r.ip or "")

    widths = [6, 20, 28, 12, 16, 18, 38, 60, 60, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_logs(
    db: AsyncSession,
    *,
    actor_user_id: UUID | None = None,
    action: str | None = None,
    entity_type: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    search: str | None = None,
    faculty_id: UUID | None = None,
    exclude_actions: list[str] | None = None,
) -> tuple[bytes, int, int]:
    """Excel: (bytes, eksport qilingan qatorlar, filtrga mos jami qatorlar)."""
    rows, total = await list_logs(
        db,
        offset=0,
        limit=EXPORT_MAX_ROWS,
        actor_user_id=actor_user_id,
        action=action,
        entity_type=entity_type,
        date_from=date_from,
        date_to=date_to,
        search=search,
        faculty_id=faculty_id,
        exclude_actions=exclude_actions,
    )
    return build_audit_xlsx(rows), len(rows), total
