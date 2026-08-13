"""ContractTemplateDoc service — DOCX shablon yuklash, placeholder aniqlash, render.

docxtpl ({{ jinja }} sintaksisi) bilan ishlaydi. Output — to'ldirilgan .docx
(QR rasm sifatida joylashtirilishi mumkin). LibreOffice talab qilinmaydi.

Yangi: `variables` JSONB — har bir placeholder uchun to'liq metadata:
  key, label, type, required, source, options, placeholder, defaultValue
"""

from __future__ import annotations

import io
import secrets
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID

from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage
from fastapi import HTTPException, UploadFile, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.contract_template import ContractTemplateDoc
from app.models.user import User
from app.schemas.contract_template import TemplateFormField, TemplateFormResponse

_STORAGE = Path(__file__).parent.parent.parent / "storage" / "contract_templates"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MAX_SIZE = 10 * 1024 * 1024

# ─── Tizimdan avtomatik olinadigan o'zgaruvchilar ─────────────────
SYSTEM_VARIABLES = {
    "contract_number",
    "contract_no",
    "contract_day",
    "contract_month",
    "contract_year",
    "contract_date",
    "student_full_name",
    "student_first_name",
    "student_last_name",
    "faculty_name",
    "specialty_name",
    "speciality_name",
    "course",
    "group_name",
    "education_form",
    "academic_year",
    "university_name",
    "university_rector_full_name",
    "qr_code",
    "qr",
    # Eski alias'lar (backward compat)
    "fish",
    "ism",
    "student_name",
    "student_field",
    "yonalish",
    "shifr",
    "kurs",
    "guruh",
    "talim_shakli",
    "day",
    "month",
    "year",
}


def _detect_placeholders(docx_bytes: bytes) -> list[str]:
    """DOCX ichidagi {{ maydon }} placeholder'larini aniqlaydi."""
    tpl = DocxTemplate(io.BytesIO(docx_bytes))
    return sorted(tpl.get_undeclared_template_variables())


def _auto_variables_from_placeholders(placeholders: list[str]) -> list[dict[str, Any]]:
    """Eski placeholders ro'yxatidan avtomatik variables metadata hosil qiladi.

    Agar shablon faqat eski format (placeholders: ["fish", "obyekt"]) da bo'lsa,
    har bir placeholder uchun source'ni aniqlaydi.
    """
    variables = []
    for key in placeholders:
        source = "system" if key in SYSTEM_VARIABLES else "student_input"
        label = key.replace("_", " ").title()
        variables.append(
            {
                "key": key,
                "label": label,
                "type": "text",
                "required": True,
                "source": source,
            }
        )
    return variables


def _save_docx(content: bytes, original: str) -> dict[str, Any]:
    now = datetime.now(UTC)
    _STORAGE.mkdir(parents=True, exist_ok=True)
    rand = secrets.token_urlsafe(12).replace("-", "").replace("_", "")[:16]
    filename = f"{rand}.docx"
    (_STORAGE / filename).write_bytes(content)
    return {
        "name": original or filename,
        "path": filename,
        "mime": _DOCX_MIME,
        "size": len(content),
        "uploaded_at": now.isoformat(),
    }


def _abs_path(rel: str) -> Path:
    """Faqat shablon nomi (path traversal'siz)."""
    candidate = (_STORAGE / Path(rel).name).resolve()
    base = _STORAGE.resolve()
    try:
        candidate.relative_to(base)
    except ValueError as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Noto'g'ri yo'l") from e
    return candidate


def _get_effective_variables(tpl: ContractTemplateDoc) -> list[dict[str, Any]]:
    """Shablon uchun effektiv variablelar ro'yxatini qaytaradi.

    Agar `variables` JSONB to'ldirilgan bo'lsa — undan oladi.
    Aks holda eski `placeholders` dan avtomatik hosil qiladi.
    """
    if tpl.variables:
        return tpl.variables
    return _auto_variables_from_placeholders(tpl.placeholders or [])


async def get_form_fields(db: AsyncSession, template_id: UUID) -> TemplateFormResponse:
    """Talaba shablon tanlaganda faqat student_input fieldlarni qaytaradi.

    source=system bo'lgan variablelar frontendga yuborilmaydi — talaba ularni ko'rmaydi.
    """
    tpl = await get_template(db, template_id)
    variables = _get_effective_variables(tpl)

    fields = []
    for var in variables:
        if var.get("source") == "system":
            continue  # system variablelarni ko'rsatmaymiz
        fields.append(
            TemplateFormField(
                key=var["key"],
                label=var.get("label", var["key"].replace("_", " ").title()),
                type=var.get("type", "text"),
                required=var.get("required", True),
                options=var.get("options"),
                placeholder=var.get("placeholder"),
                defaultValue=var.get("defaultValue"),
            )
        )

    return TemplateFormResponse(
        templateId=str(tpl.id),
        templateName=tpl.name,
        fields=fields,
    )


def validate_student_input(
    tpl: ContractTemplateDoc, variable_values: dict[str, Any]
) -> dict[str, Any]:
    """Talaba yuborgan variable_values ni template schema bilan tekshiradi.

    Xavfsizlik:
    - Faqat source=student_input variablelarni qabul qiladi
    - source=system variablelarni frontenddan qabul QILMAYDI
    - Template'da yo'q variablelarni qabul QILMAYDI
    - required=True lekin bo'sh qiymat yuborilsa — xato
    - type=select bo'lsa — faqat options'dagi qiymatlarni qabul qiladi
    """
    variables = _get_effective_variables(tpl)
    allowed_keys: set[str] = set()
    variable_map: dict[str, dict[str, Any]] = {}

    for var in variables:
        if var.get("source") == "student_input":
            allowed_keys.add(var["key"])
            variable_map[var["key"]] = var

    # Faqat ruxsat berilgan keylarni olish (qo'shimcha fake variablelarni rad etish)
    validated: dict[str, Any] = {}
    for key, value in variable_values.items():
        if key not in allowed_keys:
            continue  # silently ignore unauthorized keys
        validated[key] = value

    # Required tekshiruvi
    errors: list[str] = []
    for key, var_meta in variable_map.items():
        if var_meta.get("required", True):
            val = validated.get(key)
            if val is None or (isinstance(val, str) and not val.strip()):
                label = var_meta.get("label", key)
                errors.append(f"«{label}» maydoni to'ldirilishi shart")

        # Select validation
        if var_meta.get("type") == "select" and key in validated:
            options = var_meta.get("options", [])
            if options and validated[key] not in options:
                label = var_meta.get("label", key)
                errors.append(f"«{label}» uchun noto'g'ri qiymat: {validated[key]}")

    if errors:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="; ".join(errors),
        )

    return validated


async def create_template(
    db: AsyncSession,
    *,
    file: UploadFile,
    name: str,
    description: str | None,
    practice_type_id: UUID | None,
    user: User,
) -> ContractTemplateDoc:
    if not (file.filename or "").lower().endswith(".docx"):
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, ".docx fayl yuklang")
    content = await file.read()
    if not content:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Fayl bo'sh")
    if len(content) > _MAX_SIZE:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            f"Fayl juda katta (max {_MAX_SIZE // 1024 // 1024} MB)",
        )
    try:
        placeholders = _detect_placeholders(content)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"DOCX o'qib bo'lmadi: {e}") from e

    # Placeholderlardan avtomatik variables hosil qilish
    variables = _auto_variables_from_placeholders(placeholders)

    attachment = _save_docx(content, file.filename or "shablon.docx")
    tpl = ContractTemplateDoc(
        name=name,
        description=description,
        practice_type_id=practice_type_id,
        file_attachment=attachment,
        placeholders=placeholders,
        variables=variables,
        created_by_id=user.id,
    )
    db.add(tpl)
    await db.commit()
    await db.refresh(tpl)
    return tpl


async def create_html_template(
    db: AsyncSession,
    *,
    name: str,
    html_content: str,
    description: str | None,
    practice_type_id: UUID | None,
    user: User,
    variables: list[dict[str, Any]] | None = None,
) -> ContractTemplateDoc:
    """HTML kontent bilan yangi shablon yaratish (DOCX shart emas)."""
    import re

    # {variable_name} ko'rinishidagi o'zgaruvchilarni aniqlash
    placeholders = sorted(set(re.findall(r"\{(\w+)\}", html_content)))

    # Agar frontenddan variables kelgan bo'lsa — shu ishlatiladi
    # Aks holda placeholderlardan avtomatik hosil qilinadi
    if not variables:
        variables = _auto_variables_from_placeholders(placeholders)

    tpl = ContractTemplateDoc(
        name=name,
        description=description,
        practice_type_id=practice_type_id,
        html_content=html_content,
        placeholders=placeholders,
        variables=variables,
        created_by_id=user.id,
    )
    db.add(tpl)
    await db.commit()
    await db.refresh(tpl)
    return tpl


async def list_templates(db: AsyncSession) -> list[ContractTemplateDoc]:
    return list(
        (
            await db.execute(
                select(ContractTemplateDoc).order_by(ContractTemplateDoc.created_at.desc())
            )
        )
        .scalars()
        .all()
    )


async def get_template(db: AsyncSession, id_: UUID) -> ContractTemplateDoc:
    tpl = await db.get(ContractTemplateDoc, id_)
    if not tpl:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Shablon topilmadi")
    return tpl


async def update_template(db: AsyncSession, id_: UUID, data: dict[str, Any]) -> ContractTemplateDoc:
    tpl = await get_template(db, id_)
    for key, value in data.items():
        if value is not None:
            setattr(tpl, key, value)

    # Agar html_content yangilansa, placeholderlarni qayta aniqlash
    if "html_content" in data and data["html_content"]:
        import re

        placeholders = sorted(set(re.findall(r"\{(\w+)\}", data["html_content"])))
        tpl.placeholders = placeholders
        # Agar variables yangilanmagan bo'lsa — avtomatik hosil qilish
        if "variables" not in data or not data.get("variables"):
            tpl.variables = _auto_variables_from_placeholders(placeholders)

    await db.commit()
    await db.refresh(tpl)
    return tpl


async def delete_template(db: AsyncSession, id_: UUID) -> None:
    tpl = await get_template(db, id_)
    path = (tpl.file_attachment or {}).get("path")
    await db.delete(tpl)
    await db.commit()
    if path:
        with __import__("contextlib").suppress(Exception):
            _abs_path(path).unlink(missing_ok=True)


def template_file_path(tpl: ContractTemplateDoc) -> Path:
    rel = (tpl.file_attachment or {}).get("path")
    if not rel:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Shablon fayli yo'q")
    return _abs_path(rel)


def render_docx(
    tpl: ContractTemplateDoc, context: dict[str, Any], qr_png: bytes | None = None
) -> bytes:
    """Shablonni context bilan to'ldirib, to'ldirilgan .docx baytlarini qaytaradi.

    `qr_png` berilsa va shablonda `{{ qr }}` placeholder bo'lsa — QR rasm joylashtiriladi.
    """
    doc = DocxTemplate(str(template_file_path(tpl)))
    ctx = dict(context)
    if qr_png is not None and "qr" in (tpl.placeholders or []):
        ctx["qr"] = InlineImage(doc, io.BytesIO(qr_png), width=Mm(30))
    doc.render(ctx)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
